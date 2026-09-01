from __future__ import annotations

import unittest

from PIL import Image
from docx import Document

from server.pdf_to_docx_core.models import LineBox, WordBox
from server.pdf_to_docx_core.transcript import (
    TRANSCRIPT_COURSE_TITLES,
    choose_consecutive,
    clean_ocr_token,
    cluster_numeric,
    cluster_y_centers,
    detect_transcript_grid_geometry,
    join_positioned_words,
    normalize_grade_cell,
    normalize_transcript_course_cell,
    set_cell_margins,
    set_cell_width,
    set_row_height,
    set_table_cell_text,
    set_table_layout,
    token_has_text_signal,
    transcript_row_centers,
    words_from_lines,
    words_in_box,
)


def make_word(
    text: str,
    left: float,
    top: float,
    *,
    width: float = 30,
    height: float = 12,
    conf: float = 90,
) -> WordBox:
    return WordBox(text, left, top, width, height, conf, 1, 1, 1)


def make_line(words: list[WordBox]) -> LineBox:
    left = min(word.left for word in words)
    top = min(word.top for word in words)
    right = max(word.left + word.width for word in words)
    bottom = max(word.top + word.height for word in words)
    return LineBox(
        " ".join(word.text for word in words),
        words,
        left,
        top,
        right - left,
        bottom - top,
        sum(word.conf for word in words) / len(words),
        1000,
        1400,
        500,
        700,
        [(left, right, " ".join(word.text for word in words))],
    )


class TranscriptModelTests(unittest.TestCase):
    def test_clean_token_normalizes_brackets_and_discards_rule_noise(self) -> None:
        self.assertEqual(clean_ocr_token(" [good] "), "(good)")
        self.assertEqual(clean_ocr_token("--"), "")

    def test_grade_normalization_repairs_common_ocr_substitutions(self) -> None:
        self.assertEqual(normalize_grade_cell("5/S"), "5/5")
        self.assertEqual(normalize_grade_cell("8 10"), "8/10")
        self.assertEqual(normalize_grade_cell("40/10"), "10/10")
        self.assertEqual(normalize_grade_cell("10 / 10"), "10/10")

    def test_course_normalization_uses_expected_curriculum_for_blank_or_matching_cells(
        self,
    ) -> None:
        self.assertEqual(normalize_transcript_course_cell("", 1), "State Language")
        self.assertEqual(
            normalize_transcript_course_cell("Russian Lang", 3), "Russian Language"
        )
        self.assertEqual(normalize_transcript_course_cell("987654", 2), "987654")
        self.assertEqual(len(TRANSCRIPT_COURSE_TITLES), 23)

    def test_text_signal_accepts_identifiers_and_table_markers(self) -> None:
        self.assertTrue(token_has_text_signal("A1"))
        self.assertTrue(token_has_text_signal("#"))
        self.assertFalse(token_has_text_signal("()"))

    def test_join_positioned_words_merges_tight_fragments_and_keeps_rows(self) -> None:
        words = [
            make_word("File", 10, 10, width=25),
            make_word("Mint", 40, 10, width=25),
            make_word("Second", 10, 60, width=40),
        ]
        self.assertEqual(join_positioned_words(words), "FileMint\nSecond")

    def test_words_from_lines_removes_duplicate_word_boxes(self) -> None:
        duplicate_a = make_word("same", 10, 10)
        duplicate_b = make_word("same", 11, 10)
        unique = make_word("unique", 100, 40)
        words = words_from_lines([make_line([duplicate_a, duplicate_b, unique])])
        self.assertEqual([word.text for word in words], ["same", "unique"])

    def test_numeric_clustering_uses_medians(self) -> None:
        self.assertEqual(cluster_numeric([1, 2, 3, 20, 21], 3), [2.0, 20.5])
        self.assertEqual(cluster_y_centers([10, 12, 100, 104], 10), [11, 102])

    def test_choose_consecutive_prefers_the_widest_contiguous_window(self) -> None:
        self.assertEqual(choose_consecutive([0, 1, 2, 10, 20], 3), [2, 10, 20])
        self.assertEqual(choose_consecutive([1, 2], 5), [1, 2])

    def test_words_in_box_uses_centers_and_confidence(self) -> None:
        inside = make_word("inside", 10, 10, conf=90)
        low = make_word("low", 20, 20, conf=10)
        outside = make_word("outside", 200, 200, conf=90)
        self.assertEqual(words_in_box([inside, low, outside], 0, 0, 100, 100), [inside])

    def test_transcript_row_centers_returns_stable_fallback_grid(self) -> None:
        centers = transcript_row_centers([], 1000, 1400)
        self.assertEqual(len(centers), 23)
        self.assertLess(centers[0], centers[-1])

    def test_grid_detection_rejects_blank_pages(self) -> None:
        blank = Image.new("RGB", (600, 800), "white")
        self.assertIsNone(detect_transcript_grid_geometry(blank, 300, 400))

    def test_docx_table_helpers_apply_fixed_layout_dimensions_and_content(self) -> None:
        document = Document()
        table = document.add_table(rows=1, cols=1)
        cell = table.cell(0, 0)
        set_table_cell_text(cell, "Grade\n5/5", size=9, bold=True, align="center")
        set_cell_width(cell, 72)
        set_cell_margins(cell, 3)
        set_row_height(table.rows[0], 24, "exact")
        set_table_layout(table, width_pt=72, indent_pt=10)

        xml = table._tbl.xml
        self.assertIn("Grade", cell.text)
        self.assertIn("5/5", cell.text)
        self.assertIn("tblLayout", xml)
        self.assertIn("trHeight", xml)
        self.assertIn("tcMar", xml)


if __name__ == "__main__":
    unittest.main()
