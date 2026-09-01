"""Transcript reconstruction tests for real grid geometry and row selection."""

from __future__ import annotations

import unittest

from docx import Document
from docx.oxml.ns import qn
from PIL import Image, ImageDraw

from server.pdf_to_docx_core.models import WordBox
from server.pdf_to_docx_core.transcript import (
    detect_transcript_grid_geometry,
    join_positioned_words,
    set_cell_width,
    set_row_height,
    set_table_cell_text,
    set_table_layout,
    transcript_row_centers,
)


def word(
    text: str,
    left: float,
    top: float,
    *,
    width: float = 30,
    height: float = 12,
    conf: float = 90,
) -> WordBox:
    return WordBox(text, left, top, width, height, conf, 1, 1, 1)


class TranscriptGeometryTests(unittest.TestCase):
    def test_grid_detector_recovers_six_columns_and_last_twenty_five_rows(self) -> None:
        image = Image.new("RGB", (1000, 1400), "white")
        draw = ImageDraw.Draw(image)
        for x in (150, 300, 450, 600, 750, 900):
            draw.line((x, 400, x, 1200), fill="black", width=4)
        for y in range(420, 1171, 30):
            draw.line((150, y, 900, y), fill="black", width=4)

        geometry = detect_transcript_grid_geometry(image, 500, 700)

        self.assertIsNotNone(geometry)
        assert geometry is not None
        self.assertEqual(len(geometry["xPositionsPx"]), 6)
        self.assertEqual(len(geometry["yPositionsPx"]), 25)
        self.assertEqual(len(geometry["colWidthsPt"]), 5)
        self.assertEqual(len(geometry["rowHeightsPt"]), 24)
        self.assertAlmostEqual(geometry["tableWidthPt"], 375.0, delta=3)

    def test_positioned_word_joining_handles_punctuation_noise_and_rows(self) -> None:
        self.assertEqual(join_positioned_words([]), "")
        words = [
            word("A", 10, 10, width=20),
            word("-", 31, 10, width=4),
            word("B", 36, 10, width=10),
            word("separate", 100, 10, width=50),
            word("--", 170, 10, width=20),
            word("Second", 10, 80, width=50),
        ]
        self.assertEqual(join_positioned_words(words), "A-B separate\nSecond")

    def test_row_center_detection_prefers_eighteen_observed_course_rows(self) -> None:
        observed = [
            word(f"Course{index}", 280, 700 + index * 58) for index in range(18)
        ]
        ignored = [
            word("low confidence", 280, 710, conf=20),
            word("123", 280, 770),
            word("outside", 800, 830),
        ]
        centers = transcript_row_centers(observed + ignored, 1000, 2000)
        self.assertEqual(len(centers), 18)
        self.assertAlmostEqual(centers[0], 706)
        self.assertAlmostEqual(centers[-1], 1692)

    def test_docx_helpers_update_existing_xml_dimensions_and_alignment(self) -> None:
        document = Document()
        table = document.add_table(rows=1, cols=1)
        cell = table.cell(0, 0)
        tc_pr = cell._tc.get_or_add_tcPr()
        default_width = tc_pr.find(qn("w:tcW"))
        if default_width is not None:
            tc_pr.remove(default_width)

        set_cell_width(cell, 50)
        set_cell_width(cell, 72)
        set_row_height(table.rows[0], 18)
        set_row_height(table.rows[0], 24, "exact")
        set_table_layout(table)
        set_table_layout(table, width_pt=120, indent_pt=-10)
        set_table_cell_text(cell, "Right", align="right")
        set_table_cell_text(cell, "Left", align="left")

        xml = table._tbl.xml
        self.assertIn('w:w="1440"', xml)
        self.assertIn('w:hRule="exact"', xml)
        self.assertIn('w:w="0"', xml)
        self.assertEqual(cell.text, "Left")


if __name__ == "__main__":
    unittest.main()
